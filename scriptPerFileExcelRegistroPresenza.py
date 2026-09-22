import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# --- STILI ---
title_font = Font(name='Arial', size=16, bold=True)
header_font = Font(name='Arial', size=11, bold=True)
thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                     top=Side(style='thin'), bottom=Side(style='thin'))
center_align = Alignment(horizontal='center', vertical='center')
header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

# ==========================================
# FOGLIO 1: COPERTINA
# ==========================================
ws1 = wb.active
ws1.title = "Copertina"
ws1.merge_cells('A1:E10')
cell_title = ws1['A1']
cell_title.value = "REGISTRO GENERALE PRESENZE"
cell_title.font = Font(name='Arial', size=24, bold=True)
cell_title.alignment = center_align

ws1.merge_cells('A12:E12')
ws1['A12'].value = "Ente / Azienda: _________________________"
ws1['A12'].font = Font(size=14)
ws1['A12'].alignment = center_align

# ==========================================
# FOGLIO 2: REGISTRO GENERALE
# ==========================================
ws2 = wb.create_sheet("Registro Generale")
headers_gen = ["ID", "Cognome", "Nome"] + [f"Lez {i}" for i in range(1, 11)] + ["Totale Presenze"]

for col_num, data in enumerate(headers_gen, 1):
    cell = ws2.cell(row=1, column=col_num, value=data)
    cell.font = header_font
    cell.fill = header_fill
    cell.border = thin_border
    cell.alignment = center_align

# Esempio dati e formula
for row in range(2, 20):
    ws2.cell(row=row, column=1, value=row-1).border = thin_border
    for col in range(2, 15):
        ws2.cell(row=row, column=col).border = thin_border
    
    # Formula CONTA.SE (Conta le "P" nelle colonne lezione)
    # Colonne Lezioni sono dalla 4 alla 13 (D alla M)
    formula = f'=COUNTIF(D{row}:M{row}, "P")'
    cell_tot = ws2.cell(row=row, column=14, value=formula)
    cell_tot.font = Font(bold=True)
    cell_tot.border = thin_border
    cell_tot.alignment = center_align

ws2.column_dimensions['A'].width = 5
ws2.column_dimensions['B'].width = 20
ws2.column_dimensions['C'].width = 20

# ==========================================
# FOGLIO 3: TEMPLATE LEZIONE
# ==========================================
ws3 = wb.create_sheet("Template Lezione")

# Intestazione
ws3.merge_cells('A1:F1')
ws3['A1'].value = "REGISTRO DELLA LEZIONE"
ws3['A1'].font = title_font
ws3['A1'].alignment = center_align

info_labels = ["Titolo:", "Data:", "Docente:"]
info_pos = ['B1', 'A3', 'A4'] # Semplificato per l'esempio
ws3['A3'].value = "Titolo Lezione:"; ws3['B3'].value = "______________________"
ws3['A4'].value = "Data:"; ws3['B4'].value = "___/___/___"
ws3['D4'].value = "Orario:"; ws3['E4'].value = "________"
ws3['A5'].value = "Docente:"; ws3['B5'].value = "______________________"

# Tabella Firme
headers_lesson = ["N.", "Cognome", "Nome", "Firma Entrata", "Firma Uscita", "Note"]
for col_num, data in enumerate(headers_lesson, 1):
    cell = ws3.cell(row=7, column=col_num, value=data)
    cell.font = header_font
    cell.fill = header_fill
    cell.border = thin_border
    cell.alignment = center_align

# Righe per firme (Altezza aumentata per scrittura a mano)
for row in range(8, 30):
    ws3.cell(row=row, column=1, value=row-7).alignment = center_align
    for col in range(1, 7):
        cell = ws3.cell(row=row, column=col)
        cell.border = thin_border
    ws3.row_dimensions[row].height = 30 # Spazio per firma

# Argomenti
ws3.merge_cells('A32:F32')
ws3['A32'].value = "ARGOMENTI TRATTATI:"
ws3['A32'].font = header_font

ws3.merge_cells('A33:F40')
topics_cell = ws3['A33']
topics_cell.border = thin_border
topics_cell.alignment = Alignment(wrap_text=True, vertical='top')

# Firma Docente
ws3.merge_cells('A42:C42')
ws3['A42'].value = "Firma del Docente: __________________________"

# Larghezza colonne
ws3.column_dimensions['A'].width = 5
ws3.column_dimensions['B'].width = 25
ws3.column_dimensions['C'].width = 20
ws3.column_dimensions['D'].width = 20
ws3.column_dimensions['E'].width = 20
ws3.column_dimensions['F'].width = 20

# Impostazioni di stampa
ws3.print_area = 'A1:F42'
ws3.page_setup.fitToWidth = 1
ws3.page_setup.fitToHeight = 1

# Salvataggio
file_name = "Registro_Presenze_Personale.xlsx"
wb.save(file_name)
print(f"File '{file_name}' creato con successo!")